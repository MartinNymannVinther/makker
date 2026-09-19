"""
Tekstudtræk fra vedhæftede filer.

Ét sted at slå op: "hvilke filtyper kan vi læse, og hvordan". app.py kender
kun funktionen extract() — den ved intet om PDF eller Word.

Alt hvad der kommer ud herfra har samme form (se Attachment), og feltet
`kind` fortæller hvordan indholdet skal bruges:

    kind="text"   →  teksten lægges ind i samtalen som tekst
    kind="image"  →  (endnu ikke bygget) billedet sendes som billede til en
                     model der kan se. Se stubben nederst i filen.
"""
import io
import os
from dataclasses import dataclass, asdict

from pypdf import PdfReader
from docx import Document

# Grænser. Store filer koster tokens, og en samtale skal stadig kunne sendes.
MAX_BYTES = 10 * 1024 * 1024   # 10 MB rå fil
MAX_CHARS = 60_000             # ca. 15.000 tokens udtrukket tekst


class UploadError(Exception):
    """Fejl vi kan forklare brugeren i klar tekst."""


def _antal(n: int, ental: str, flertal: str) -> str:
    """"1 side" / "12 sider" — chippen i UI'et skal læse som dansk."""
    return f"{n} {ental if n == 1 else flertal}"


@dataclass
class Attachment:
    name: str      # filnavn, vises i UI'et
    kind: str      # "text" (senere også "image")
    note: str      # kort beskrivelse til UI'et, fx "12 sider"
    text: str = ""  # udtrukket indhold når kind == "text"

    def as_dict(self):
        return asdict(self)


# --- Én udtrækker pr. filtype -------------------------------------------
# Hver funktion tager rå bytes og returnerer (tekst, kort note).

def _from_pdf(data: bytes):
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception:
        raise UploadError("PDF'en kunne ikke læses — er filen beskadiget?")

    if reader.is_encrypted:
        raise UploadError("PDF'en er låst med kodeord.")

    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n\n".join(pages), _antal(len(pages), "side", "sider")


def _from_docx(data: bytes):
    try:
        doc = Document(io.BytesIO(data))
    except Exception:
        raise UploadError("Word-filen kunne ikke læses. Er det en gammel .doc? "
                          "Gem den som .docx og prøv igen.")

    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # python-docx lister tabeller for sig, så de havner samlet til sidst.
    # Det er fint til formålet — modellen skal bare kunne læse indholdet.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts), _antal(len(parts), "afsnit", "afsnit")


def _from_plain(data: bytes):
    text = data.decode("utf-8", errors="replace")
    return text, _antal(len(text.splitlines()), "linje", "linjer")


# Vil du understøtte en filtype mere, tilføjer du én linje her.
EXTRACTORS = {
    ".pdf":  _from_pdf,
    ".docx": _from_docx,
    ".txt":  _from_plain,
    ".md":   _from_plain,
}

# Bruges af UI'et til at sætte accept="..." på filvælgeren.
SUPPORTED = sorted(EXTRACTORS)


def extract(filename: str, data: bytes) -> Attachment:
    """Rå fil ind, ensartet Attachment ud. Kaster UploadError ved problemer."""
    if not data:
        raise UploadError("Filen er tom.")
    if len(data) > MAX_BYTES:
        raise UploadError(f"Filen er for stor (max {MAX_BYTES // 1024 // 1024} MB).")

    name = os.path.basename(filename or "fil")
    ext = os.path.splitext(name)[1].lower()

    reader = EXTRACTORS.get(ext)
    if reader is None:
        raise UploadError(f"{ext or 'Filtypen'} understøttes ikke endnu. "
                          f"Prøv: {', '.join(SUPPORTED)}.")

    text, note = reader(data)
    text = text.strip()

    if not text:
        raise UploadError("Der blev ikke fundet tekst i filen. "
                          "Er det en scanning eller et billede?")

    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + "\n\n[… resten af filen er klippet fra]"
        note += " · forkortet"

    return Attachment(name=name, kind="text", note=note, text=text)


# --- Næste skridt: billeder ----------------------------------------------
# Når en model der kan se skal have billedet med, ser det sådan ud:
#
#   def _from_image(data):  →  Attachment(kind="image", data_url="data:image/png;base64,...")
#
# Attachment får et felt mere, EXTRACTORS får ".png"/".jpg", og app.py
# pakker kind="image" som en content-blok i stedet for som tekst:
#
#   {"role": "user", "content": [
#       {"type": "image_url", "image_url": {"url": att["data_url"]}},
#       {"type": "text", "text": spørgsmålet},
#   ]}
#
# Resten af kæden — upload, chip i UI'et, historik — er den samme.
